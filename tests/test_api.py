"""Prueba de extremo a extremo de la API y de la lógica de coberturas.

Uso:  DATA_DIR=/tmp/supply-test python tests/test_api.py ruta/al/MM_Supply.xlsx

DATA_DIR debe ser una carpeta vacía o nueva: la prueba crea su propia base de datos.
Si hay Node.js instalado, comprueba también que static/core.js evalúa igual que core.py.
"""
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile

os.environ.setdefault("ADMIN_USER", "admin")
os.environ.setdefault("ADMIN_PASSWORD", "admin12345")
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "app")
sys.path.insert(0, ROOT)
import app as A  # noqa: E402
import core  # noqa: E402

H = {"X-Requested-With": "supply-app"}
data = open(sys.argv[1], "rb").read()


def upload(c, dry, fecha=""):
    return c.post("/api/upload", data={"file": (io.BytesIO(data), "MM_Supply.xlsx"), "dry": "1" if dry else "0", "fecha": fecha},
                  headers=H, content_type="multipart/form-data")


def check(name, cond, detail=""):
    print(("OK   " if cond else "FALLO"), name, detail)
    if not cond:
        check.fails += 1


def check_caso(name, casos, cond, detail=""):
    """Como check, para comprobaciones que necesitan un caso concreto en el fichero: si no lo trae, se avisa y no falla."""
    if casos:
        check(f"{name} ({casos} casos)", cond, detail)
    else:
        print("--   ", name, "(sin caso en este fichero)")


check.fails = 0
c = A.app.test_client()

# Sesión y protección de peticiones
check("sin cabecera de la app se rechaza", c.post("/api/login", json={"username": "admin", "password": "admin12345"}).status_code == 403)
check("login del administrador", c.post("/api/login", json={"username": "admin", "password": "admin12345"}, headers=H).status_code == 200)

# Carga de datos
r = upload(c, True)
check("comprobación del fichero", r.status_code == 200, r.json.get("preview") if r.json else r.status_code)
check("publicación con datos de la semana anterior", upload(c, False, "2026-09-22").json.get("ok"))
check("publicación con datos de hoy", upload(c, False).json.get("ok"))
r = upload(c, True)
check("la comprobación avisa de que sustituye la carga del mismo día", bool(r.json["preview"].get("sustituye")), r.json["preview"].get("sustituye"))
check("republicar el mismo día sustituye la carga", upload(c, False).json.get("ok") and len(c.get("/api/loads").json) == 2, len(c.get("/api/loads").json))
ds = c.get("/api/dataset").json
check("referencias publicadas", len(ds["refs"]) > 0, f'{len(ds["refs"])} · previsión {ds["meta"]["version"]}')
# Versiones de la previsión por referencia: aparte del dataset; la vigente cuadra con la previsión operativa de sus meses
check("el dataset no lleva las versiones de previsión", "pvers" not in ds)
_rv = next(r for r in ds["refs"] if sum(r["pv"]) > 0)
_V = c.get(f'/api/desviacion/{_rv["k"]}/versiones').json
_src = ds["meta"]["prev_src"]
check("versiones: guardadas y con la vigente", not _V["sin"] and ds["meta"]["version"] in _V["p"], (_V.get("sin"), list(_V.get("p", {}))))
check("versiones: la vigente cuadra con la previsión operativa", all(round(_V["p"][v].get(str(i), 0)) == _rv["pv"][i] for i, v in enumerate(_src) if v),
      [(i, v, _V["p"].get(v, {}).get(str(i)), _rv["pv"][i]) for i, v in enumerate(_src) if v][:4])
check("versiones: venta del mes en curso", _V["s"].get("0", 0) == _rv["v0"], (_V["s"].get("0"), _rv["v0"]))
check("versiones: referencia no válida", c.get("/api/desviacion/" + "9" * 25 + "/versiones").status_code == 400)
check("versiones: referencia fuera de la carga", c.get("/api/desviacion/999999999999/versiones").status_code == 404)
# Capacidad llega hasta el último mes de la versión vigente: meses extra con su previsión (px), pedidos (pdx) y OF (ox)
_hx = ds["meta"].get("hx")
check("meses de previsión más allá de los 12 (capacidad)", isinstance(_hx, int) and 0 <= _hx <= 12
      and all(len(r["px"]) == _hx for r in ds["refs"] if "px" in r) and (_hx == 0 or any("px" in r for r in ds["refs"]))
      and all(_hx and 12 <= m < 12 + _hx for r in ds["refs"] for m, _, _ in r.get("ox", [])), _hx)
# Propuestas sin fijar que ya no llegan con el plazo de fabricación: a la primera fecha posible (hoy + 21 días)
import datetime as _dt  # noqa: E402
_lim = (_dt.date.fromisoformat(ds["meta"]["hoy"]) + _dt.timedelta(days=21)).isoformat()
_ps = [e for r in ds["refs"] for e in r["en"] if e["t"] == "P"]
check("propuestas sin fijar no antes de hoy + 21 días", _ps and all(e["d"] >= _lim for e in _ps), [e for e in _ps if e["d"] < _lim][:2])
_ao = [p for r in ds["refs"] for p in r.get("ao", [])]
_lim30 = (_dt.date.fromisoformat(ds["meta"]["hoy"]) - _dt.timedelta(days=30)).isoformat()
check("atrasados de más de 30 días fuera de la demanda y listados", _ao and all(p[0] < _lim30 for p in _ao) and all(x[0] >= _lim30 for r in ds["refs"] for x in r["pdd"]), len(_ao))
check("propuestas movidas guardan su fecha del MRP", any(e.get("dm") for e in _ps) and all(e["dm"] < e["d"] == _lim for e in _ps if e.get("dm")))
pf = ds.get("porfolio") or {}
_hoy = _dt.date.fromisoformat(ds["meta"]["hoy"])
_lanz = core._hace_meses(_hoy, core.LANZ_MESES).isoformat()  # altas desde aquí: lanzamientos
check("porfolio: resumen del maestro", pf.get("res", {}).get("seguimiento") == len(ds["refs"]) and pf["res"]["fuera"] == len(pf["fuera"]), pf.get("res"))
check("porfolio: fuera solo quedan activos sin movimiento que no son lanzamientos",
      pf.get("fuera") is not None and all(x["cat"] != "prev_futura" and x["alta"] < _lanz for x in pf["fuera"]),
      [(x["k"], x["cat"], x["alta"]) for x in pf.get("fuera", []) if x["cat"] == "prev_futura" or x["alta"] >= _lanz][:3])
check("porfolio: inactivos con stock, todos con stock > 0", all(x["st"] > 0 for x in pf.get("inact", [])), len(pf.get("inact", [])))
check("cada referencia lleva su fecha de alta", sum(bool(r.get("al")) for r in ds["refs"]) > len(ds["refs"]) * 0.9, sum(bool(r.get("al")) for r in ds["refs"]))
activos = {r["k"] for r in ds["refs"]} | {x["k"] for x in pf.get("fuera", [])}
check("el sucesor de la ficha existe y está activo", all(r["sc"] in activos for r in ds["refs"] if r.get("sc")), [(r["k"], r["sc"]) for r in ds["refs"] if r.get("sc") and r["sc"] not in activos][:3])
check("los lanzamientos siempre están en seguimiento", pf.get("lanz") and all(x["app"] for x in pf["lanz"]), sum(not x["app"] for x in pf.get("lanz", [])))
check("porfolio: lanzamientos de los últimos 9 meses", len(pf.get("lanz", [])) > 0 and all(x["alta"] >= _lanz for x in pf["lanz"]), len(pf.get("lanz", [])))
# Mismo Excel con otra fecha: solo puede cambiar lo que depende de la ventana de venta (el mes que se cierra entre
# las dos fechas), nunca referencias con stock, previsión, pedidos o entradas
_cb = pf.get("cambios") or {}
_mov = {r["k"] for r in ds["refs"] if r["st"] or any(r["pv"]) or any(r["pd"]) or r["en"]}
check("porfolio: mismo Excel que la semana anterior, solo cambia la ventana de venta",
      _cb.get("salen") is not None and not any(x["k"] in _mov for x in _cb["entran"]) and all(x["m"] == "vuelve" for x in _cb["entran"])
      and all(x["m"] == "sin_mov" for x in _cb["salen"]), _cb)
r = upload(c, True)
check("la comprobación del fichero resume altas y bajas", r.json["preview"].get("cambios") == {"entran": len(_cb.get("entran", [])), "salen": len(_cb.get("salen", []))}, r.json["preview"].get("cambios"))
bel = [r for r in ds["refs"] if r["md"] == "Belloch" and r["gp"] == "Contra Stock"]
check("stock mínimo de Belloch desde mindest del maestro", sum(r["mn"] > 0 for r in bel) > len(bel) // 3, f'{sum(r["mn"] > 0 for r in bel)} de {len(bel)}')
# Pedidos de compra a proveedor (MM_PedCompras): entrada «PC»; los de LABORATORIOS BELLOCH (intragrupo) no cuentan
# Lo esperado se saca del propio fichero, leído aparte: las pruebas valen con cualquier MM_Supply
RW = core.read_workbook(sys.argv[1])
_g, _c, _n, _d = core._get, core._code, core._num, core._date
kref = {r["k"] for r in ds["refs"]}
pcs = {r["k"]: [e for e in r["en"] if e["t"] == "PC"] for r in ds["refs"]}
_got = sorted((k, e["id"], e["q"], e["d"]) for k, es in pcs.items() for e in es)
_esp, _intra = [], []
if RW.get("MM_PedCompras"):
    tc = core._Table(RW["MM_PedCompras"], "Articulo", "MM_PedCompras")
    jc, jq, jf, jp, jv = tc.col("Articulo"), tc.col("SumCantidad_Pendiente"), tc.col("Fecha Entrega"), tc.opt("Nº Pedido"), tc.opt("Nombre proveedor")
    for x in tc.data:
        k, q, d = _c(_g(x, jc)), _n(_g(x, jq)), _d(_g(x, jf))
        if k not in kref or q <= 0 or (d and (d.year - _hoy.year) * 12 + d.month - _hoy.month >= 12):
            continue
        fila = (k, core._id(_g(x, jp)), round(q), d.isoformat() if d else "")
        (_intra if "LABORATORIOS BELLOCH" in core._norm(_g(x, jv)).upper() else _esp).append(fila)
_falta = [x for x in _esp if x not in _got] + [x for x in _got if x not in _esp]
check_caso("pedidos de compra de PT leídos tal cual del fichero (12 meses)", len(_esp), not _falta, _falta[:3])
check_caso("pedidos intragrupo (LABORATORIOS BELLOCH) excluidos", len(_intra), not set(_intra) & set(_got), [x for x in _intra if x in _got][:3])
_pc = [e for es in pcs.values() for e in es]
check_caso("pedido de compra con fecha pasada marcado atrasado", sum(bool(e["d"]) and e["d"] < ds["meta"]["hoy"] for e in _pc),
           all(e["late"] == (bool(e["d"]) and e["d"] < ds["meta"]["hoy"]) for e in _pc), [e for e in _pc if e["late"] != (bool(e["d"]) and e["d"] < ds["meta"]["hoy"])][:2])
# ZT del escandallo de los PT fabricados fuera: componente del PT en Escandallos, con su QNivelPT y su stock de MM_Stocks
_ts = core._Table(RW["MM_Stocks"], "Nº Articulo", "MM_Stocks")
_st = {}
for x in _ts.data:
    _st[_c(_g(x, _ts.col("Nº Articulo")))] = _st.get(_c(_g(x, _ts.col("Nº Articulo"))), 0.0) + _n(_g(x, _ts.col("SumStock_actual")))
_bom = {}
if RW.get("Escandallos"):
    tb = core._Table(RW["Escandallos"], "Nº Articulo", "Escandallos")
    for x in tb.data:
        _bom[(_c(_g(x, tb.col("Nº Articulo"))), _c(_g(x, tb.col("NivelPT"))))] = _n(_g(x, tb.col("QNivelPT"))) or 1.0
_zts = [(r["k"], z) for r in ds["refs"] for z in r.get("zt", [])]
check_caso("ZT del escandallo en el PT fabricado fuera (cantidad por PT y stock)", len(_zts),
           all(_bom.get((k, z["k"])) == z["q"] and z["st"] == round(_st.get(z["k"], 0.0)) for k, z in _zts),
           [(k, z["k"], z["q"], z["st"]) for k, z in _zts if _bom.get((k, z["k"])) != z["q"] or z["st"] != round(_st.get(z["k"], 0.0))][:2])
check("solo los PT con pedido de compra llevan ZT", all(any(e["t"] == "PC" for e in r["en"]) for r in ds["refs"] if r.get("zt")), sum(1 for r in ds["refs"] if r.get("zt")))
# Falta ZT: lo firme del ZT (stock + OF) ÷ ZT por PT no llega a los pedidos de compra pendientes
_fz = [r for r in ds["refs"] if r.get("zt") and any((max(z["st"], 0) + sum(e["q"] for e in z["en"] if e["t"] == "OF")) / (z["q"] or 1)
                                                     < sum(e["q"] for e in r["en"] if e["t"] == "PC") - 0.5 for z in r["zt"])]
check_caso("falta ZT para el pedido", len(_fz), all(core.falta_zt(r) == "Falta ZT para el pedido" for r in _fz), [(r["k"], core.falta_zt(r)) for r in _fz][:2])
_zr = next((r for r in ds["refs"] if r.get("zt")), None)
if _zr:  # caso fabricado con una referencia real: su ZT sin stock ni OF no cubre los pedidos de compra
    _zr2 = json.loads(json.dumps(_zr))
    for z in _zr2["zt"]:
        z["st"], z["en"] = 0, [e for e in z["en"] if e["t"] != "OF"]
    check("falta ZT para el pedido (ZT sin stock ni OF)", core.falta_zt(_zr2) == "Falta ZT para el pedido", (_zr["k"], core.falta_zt(_zr2)))
check("ZT suficiente: sin aviso de falta",all(core.falta_zt(r) != "Falta ZT para el pedido" for r in ds["refs"] if r.get("zt") and r not in _fz))
# Stock mínimo de Belloch = mindest del maestro
_ta = core._Table(RW["MM_Art"], "Nº Artículo", "MM_Art")
if _ta.opt("mindest") is not None:
    _min = {_c(_g(x, _ta.col("Nº Artículo"))): round(_n(_g(x, _ta.col("mindest")))) for x in _ta.data}
    check("stock mínimo de Belloch igual al mindest del maestro", all(r["mn"] == _min.get(r["k"], 0) for r in bel),
          [(r["k"], r["mn"], _min.get(r["k"])) for r in bel if r["mn"] != _min.get(r["k"], 0)][:3])
check("previsión operativa: todos los meses tienen versión", all(ds["meta"]["prev_src"]), ds["meta"]["prev_src"])
check("se compara con la carga de la fecha anterior", ds["prev"] and ds["prev"].get("hoy") == "2026-09-22", ds["prev"] and ds["prev"].get("hoy"))
check("carga anterior evaluada en las seis combinaciones", ds["prev"] and set(ds["prev"]["sem"]) == {"OF", "OFPF", "ALL", "OF_C", "OFPF_C", "ALL_C"})

# Parámetros del ABC desde la app
cfg = c.get("/api/me").json["config"]
check("parámetros del ABC por defecto", cfg.get("abc", {}).get("cortes") == [45, 80, 95] and cfg["abc"]["freq"]["Yunsey"] == [6, 4, 2, 1], cfg.get("abc"))
nA = lambda d: sum(r["abc"] == "A" for r in d["refs"])  # noqa: E731
antes = nA(c.get("/api/dataset").json)
r = c.put("/api/config", json={"abc": dict(cfg["abc"], cortes=[60, 90, 99])}, headers=H)
check("guardar cortes nuevos", r.status_code == 200 and r.json["abc"]["cortes"] == [60, 90, 99], r.json)
check("al cambiar los cortes se recalcula el ABC de la carga vigente", nA(c.get("/api/dataset").json) > antes, (antes, nA(c.get("/api/dataset").json)))
check("cortes no válidos se rechazan", c.put("/api/config", json={"abc": dict(cfg["abc"], cortes=[80, 45, 95])}, headers=H).status_code == 400)
check("XYZ por defecto 50/100 y cada referencia con su letra", cfg["abc"].get("xyz") == [50, 100]
      and all(r.get("xyz") in ("X", "Y", "Z", "") for r in c.get("/api/dataset").json["refs"]), cfg["abc"].get("xyz"))
for malo in ([100, 50], [0, 100], [50], [50, 600], "50,100", [True, 100]):
    check(f"cortes XYZ no válidos se rechazan: {malo}", c.put("/api/config", json={"abc": dict(cfg["abc"], xyz=malo)}, headers=H).status_code == 400)
_nx = lambda: sum(r.get("xyz") == "X" for r in c.get("/api/dataset").json["refs"])  # noqa: E731
_x0 = _nx()
c.put("/api/config", json={"abc": dict(cfg["abc"], xyz=[90, 150])}, headers=H)
check("al cambiar los cortes XYZ se recalcula la carga vigente", _nx() > _x0, (_x0, _nx()))
c.put("/api/config", json={"abc": cfg["abc"]}, headers=H)
check("frecuencia no válida se rechaza", c.put("/api/config", json={"abc": dict(cfg["abc"], freq={"Belloch": [0, 6, 4, 2], "Yunsey": [6, 4, 2, 1]})}, headers=H).status_code == 400)
c.put("/api/config", json={"abc": cfg["abc"]}, headers=H)
check("el horizonte se sigue guardando aparte", c.put("/api/config", json={"horizonte": 3}, headers=H).json.get("abc", {}).get("cortes") == [45, 80, 95])

# Meses de exceso por mandante
check("exceso por defecto 6/12", cfg.get("exceso") == {"Belloch": 6, "Yunsey": 12}, cfg.get("exceso"))
r = c.put("/api/config", json={"exceso": {"Belloch": 4, "Yunsey": 9}}, headers=H)
check("guardar meses de exceso", r.status_code == 200 and r.json["exceso"] == {"Belloch": 4, "Yunsey": 9}, r.json)
for malo in ({"Belloch": 0, "Yunsey": 12}, {"Belloch": 13, "Yunsey": 12}, {"Belloch": 6.5, "Yunsey": 12}, {"Belloch": 6}, "6", {"Belloch": True, "Yunsey": 12}):
    check(f"exceso no válido se rechaza: {malo}", c.put("/api/config", json={"exceso": malo}, headers=H).status_code == 400)
nexc = lambda: sum(v == "exceso" for v in c.get("/api/dataset").json["prev"]["sem"]["ALL"].values())  # noqa: E731
c.put("/api/config", json={"exceso": {"Belloch": 1, "Yunsey": 1}}, headers=H)
n1 = nexc()  # con 1 mes hay más exceso que con 6/12, salvo que todas tengan stock máximo (mínimo + lote)
c.put("/api/config", json={"exceso": {"Belloch": 6, "Yunsey": 12}}, headers=H)
check_caso("la carga anterior se evalúa con los meses de exceso guardados", int(n1 != nexc()), nexc() < n1, (n1, nexc()))

# Validaciones de acciones
r = c.post("/api/actions", json={"ref": "<img src=x onerror=alert(1)>", "text": "x"}, headers=H)
check("acción con referencia no válida se rechaza", r.status_code == 400)
k = ds["refs"][0]["k"]
r = c.post("/api/actions", json={"ref": k, "text": "Revisar", "due": "2026-10-15"}, headers=H)
aid = r.json[0]["id"] if r.status_code == 200 else None
check("acción válida", aid is not None)
check("fecha límite no válida se rechaza", c.patch(f"/api/actions/{aid}", json={"due": "no-es-fecha"}, headers=H).status_code == 400)

# Contraseña temporal: sin cambiarla no se ven datos
pw = c.post("/api/users", json={"username": "lec", "name": "Lector", "role": "lector"}, headers=H).json["password"]
c2 = A.app.test_client()
c2.post("/api/login", json={"username": "lec", "password": pw}, headers=H)
check("con contraseña temporal /api/me responde", c2.get("/api/me").status_code == 200)
check("con contraseña temporal no hay datos", c2.get("/api/dataset").status_code == 403)
c2.post("/api/me/password", json={"current": pw, "new": "nueva-clave-1"}, headers=H)
check("tras cambiarla sí hay datos", c2.get("/api/dataset").status_code == 200)
check("un lector no puede crear acciones", c2.post("/api/actions", json={"ref": k, "text": "x"}, headers=H).status_code == 403)

# Notas: las borra quien las escribió o un administrador
nid = c.post(f"/api/notes/{k}", json={"text": "nota de prueba"}, headers=H).json[0]["id"]
check("la nota trae su autor", c.get(f"/api/notes/{k}").json[0].get("uid") is not None)
check("un lector no puede borrar notas", c2.delete(f"/api/notes/{k}/{nid}", headers=H).status_code == 403)
pw3 = c.post("/api/users", json={"username": "pla", "name": "Planificador", "role": "planificador"}, headers=H).json["password"]
c3 = A.app.test_client()
c3.post("/api/login", json={"username": "pla", "password": pw3}, headers=H)
c3.post("/api/me/password", json={"current": pw3, "new": "nueva-clave-3"}, headers=H)
check("un planificador no borra notas de otro", c3.delete(f"/api/notes/{k}/{nid}", headers=H).status_code == 403)
nid3 = c3.post(f"/api/notes/{k}", json={"text": "mía"}, headers=H).json[0]["id"]
check("un planificador borra sus notas", c3.delete(f"/api/notes/{k}/{nid3}", headers=H).status_code == 200)
r = c.delete(f"/api/notes/{k}/{nid}", headers=H)
check("el administrador borra cualquier nota", r.status_code == 200 and all(n["id"] != nid for n in r.json), r.status_code)
check("borrar una nota que no existe", c.delete(f"/api/notes/{k}/{nid}", headers=H).status_code == 404)
check("el índice de notas se actualiza", c.get("/api/notes-index").json.get(k) is None)

# Parámetros de planificación: stock mínimo, lote y stock máximo
check("niveles de servicio por defecto", c.get("/api/me").json["config"].get("ns") == {"Belloch": [95, 90, 85, 85], "Yunsey": [95, 90, 85, 85]})
check("nivel de servicio no válido se rechaza", c.put("/api/config", json={"ns": {"Belloch": [100, 90, 85, 85], "Yunsey": [95, 90, 85, 85]}}, headers=H).status_code == 400)
check("guardar niveles de servicio", c.put("/api/config", json={"ns": {"Belloch": [95, 90, 85, 85], "Yunsey": [95, 90, 85, 85]}}, headers=H).status_code == 200)
pr_ = c.get("/api/parametros").json
rows = {p["k"]: p for p in pr_["rows"]}
check("parámetros: filas solo contra stock con ABC", rows and all(p["abc"] in "ABCD" for p in rows.values()), len(rows))
check("parámetros: resumen por mandante y clase", len(pr_["resumen"]["grupos"]) == 8, [(g["md"], g["abc"]) for g in pr_["resumen"]["grupos"]])
cam = [p for p in rows.values() if p["estado"] == "cambio" and p["est"] is not None]
igu = next(p for p in rows.values() if p["estado"] in ("cambio", "igual") and p is not cam[0])
k1, k2 = cam[0]["k"], igu["k"]
dec = lambda items, motivo="", cl=c: cl.post("/api/parametros/decisiones", json={"items": items, "motivo": motivo}, headers=H)  # noqa: E731
check("decidir: manual sin motivo se rechaza", dec([{"ref": k1, "ss": {"src": "manual", "v": 1234}, "lote": {"src": "erp"}}]).status_code == 400)
check("decidir: fuente no válida se rechaza", dec([{"ref": k1, "ss": {"src": "otra"}, "lote": {"src": "erp"}}]).status_code == 400)
check("decidir: referencia fuera de ámbito se rechaza", dec([{"ref": "999999999999", "ss": {"src": "erp"}, "lote": {"src": "erp"}}]).status_code == 400)
check("decidir: un lector no puede", dec([{"ref": k1, "ss": {"src": "erp"}, "lote": {"src": "erp"}}], cl=c2).status_code == 403)
r = dec([{"ref": k1, "ss": {"src": "manual", "v": cam[0]["mn"] + 700}, "lote": {"src": "calculado"}}], "Prueba de decisión")
check("decidir: manual con motivo", r.status_code == 200 and r.json["n"] == 1, r.json)
p1 = c.get(f"/api/parametros?ref={k1}").json["rows"][0]
check("decisión vigente: decidida con sus valores", p1["estado"] == "decidido" and p1["ss"] == cam[0]["mn"] + 700 and p1["lote"] == cam[0]["lc"], p1)
check("stock máximo con lo decidido en el dataset", next(x for x in c.get("/api/dataset").json["refs"] if x["k"] == k1).get("sx") == p1["ss"] + p1["lote"])
r = dec([{"ref": k2, "ss": {"src": "erp"}, "lote": {"src": "erp"}}], "Mantener")
check("mantener el ERP: aplicada al momento", c.get(f"/api/parametros?ref={k2}").json["rows"][0]["estado"] == "aplicado")
csv = c.get("/api/parametros/abas.csv")
txt = csv.data.decode("latin-1")
check("CSV para ABAS: cabecera y formato", csv.status_code == 200 and txt.startswith("Referencia;Mandante;Stock mínimo;Lote\r\n") and "attachment" in csv.headers.get("Content-Disposition", ""), txt[:80])
check("CSV para ABAS: trae lo decidido", f'{k1};{p1["md"]};{p1["ss"]};{p1["lote"]}\r\n' in txt)
check("CSV para ABAS: no trae lo aplicado", k2 not in txt)
import openpyxl  # noqa: E402
xl = c.get("/api/parametros.xlsx")
ws = openpyxl.load_workbook(io.BytesIO(xl.data)).active
_hd = [c_.value for c_ in ws[1]]
xr = {r[0]: dict(zip(_hd, r)) for r in ws.iter_rows(min_row=2, values_only=True)}  # por cabecera, no por posición
check("Excel de parámetros: todas las filas, referencia como texto", xl.status_code == 200 and ws.cell(1, 1).value == "Referencia"
      and len(xr) == len(pr_["rows"]) and isinstance(ws.cell(2, 1).value, str), (xl.status_code, len(xr), len(pr_["rows"])))
check("Excel de parámetros: decidido y estado", xr[k1]["SM decidido"] == p1["ss"] and xr[k1]["Estado"] == "Pendiente de ABAS", xr.get(k1))
check("Excel de parámetros: lleva el XYZ", xr[k1]["XYZ"] in ("X", "Y", "Z", "—"), xr[k1].get("XYZ"))
xf = openpyxl.load_workbook(io.BytesIO(c.get("/api/parametros.xlsx?estado=decidido").data)).active
check("Excel de parámetros: respeta el filtro de estado", [r[0] for r in xf.iter_rows(min_row=2, values_only=True)] == [k1])
h = c.get(f"/api/parametros/{k1}/historial").json
check("historial con autor, motivo y valores de antes", len(h) == 1 and h[0]["motivo"] == "Prueba de decisión" and h[0]["ss_antes"] == cam[0]["mn"] and h[0]["by"], h)
check("reparto: % por defecto", c.get("/api/me").json["config"].get("reparto") == {"pct": 50})
check("reparto: % no válido se rechaza", c.put("/api/config", json={"reparto": {"pct": 150}}, headers=H).status_code == 400)
F0 = pr_["resumen"]["frec"]
check("reparto: guardar el %", c.put("/api/config", json={"reparto": {"pct": 0}}, headers=H).status_code == 200)
F1 = c.get("/api/parametros").json["resumen"]["frec"]
check("reparto: con 0 % el stock mínimo repartido no pasa del ERP (± redondeo)", F1["ss_rep"] <= F1["ss_erp"] * 1.02 and F1["ss_rep"] <= F0["ss_rep"], (F1["ss_erp"], F1["ss_rep"], F0["ss_rep"]))
k3 = next((p["k"] for p in rows.values() if p["lo"] is not None and p["k"] not in (k1, k2)), None)
if k3:
    r = dec([{"ref": k3, "ss": {"src": "reparto"}, "lote": {"src": "optimo"}}])
    p3 = c.get(f"/api/parametros?ref={k3}").json["rows"][0]
    check("decidir con reparto y frecuencia óptima", r.status_code == 200 and p3["ss"] == p3["sr"] and p3["lote"] == p3["lo"], (r.json, p3["ss"], p3["sr"], p3["lote"], p3["lo"]))
k4 = next((p["k"] for p in rows.values() if p["lo"] is None), None)
if k4:
    check("decidir con frecuencia óptima sin cifra se rechaza", dec([{"ref": k4, "ss": {"src": "erp"}, "lote": {"src": "optimo"}}]).status_code == 400)
c.put("/api/config", json={"reparto": {"pct": 50}}, headers=H)
check("plazo extra de una referencia fuera de ámbito se rechaza", c.put("/api/parametros/999999999999/plazo", json={"dias": 5}, headers=H).status_code == 400)
check("plazo extra fuera de rango se rechaza", c.put(f"/api/parametros/{k1}/plazo", json={"dias": 300}, headers=H).status_code == 400)
check("plazo extra", c.put(f"/api/parametros/{k1}/plazo", json={"dias": 21, "motivo": "Tubos 66 días"}, headers=H).status_code == 200
      and c.get(f"/api/parametros?ref={k1}").json["rows"][0]["dx"] == 21)
# Aplicado al publicar: se fuerza una decisión igual al ERP sin marcar y se republica
import sqlite3  # noqa: E402
con = sqlite3.connect(A.DB_PATH)
k3 = cam[1]["k"]
con.execute("INSERT INTO param_dec(ref,md,ss,lote,src_ss,src_lote,ss_antes,lote_antes,motivo,user_id,created) VALUES(?,?,?,?,'manual','manual',?,?,'x',1,'2026-10-01T10:00:00')",
            (k3, cam[1]["md"], cam[1]["mn"], cam[1]["lt"], cam[1]["mn"], cam[1]["lt"]))
con.commit()
check("antes de publicar: decidida", c.get(f"/api/parametros?ref={k3}").json["rows"][0]["estado"] == "decidido")
upload(c, False)
check("al publicar se marca como aplicada", c.get(f"/api/parametros?ref={k3}").json["rows"][0]["estado"] == "aplicado")
# Aplicada y cambiada después en ABAS: se ignora
con.execute("UPDATE param_dec SET ss=ss+500 WHERE ref=?", (k3,))
con.commit()
p3 = c.get(f"/api/parametros?ref={k3}").json["rows"][0]
check("aplicada y cambiada en ABAS: se ignora", p3["estado"] != "aplicado" and not p3["d"] and k3 not in c.get("/api/parametros/abas.csv").data.decode("latin-1"), p3["estado"])
con.close()

# Desviación de previsiones
dv = c.get("/api/desviacion").json
check("desviación: versión vigente y acierto por versión", dv.get("version") == ds["meta"]["version"] and dv.get("vers") and all(x["e"] >= 0 for x in dv["vers"]), (dv.get("version"), len(dv.get("vers") or [])))
check("desviación: resumen con total primero", dv["marcas"][0]["mc"] == "Total" and dv["marcas"][0]["n"] == len(dv["rows"]), dv["marcas"][0])
prop = [p for p in dv["rows"] if p["tipo"] == "propuesta"]
k1, k2 = prop[0]["k"], prop[1]["k"]
acu = lambda items, motivo="", cl=c: cl.post("/api/desviacion/acuerdos", json={"items": items, "motivo": motivo}, headers=H)  # noqa: E731
check("acuerdo manual sin motivo se rechaza", acu([{"ref": k1, "src": "manual", "pct": -15}]).status_code == 400)
check("acuerdo manual fuera de rango se rechaza", acu([{"ref": k1, "src": "manual", "pct": -95}], "x").status_code == 400)
check("acuerdo fuera de ámbito se rechaza", acu([{"ref": "999999999999", "src": "mantener"}]).status_code == 400)
check("un lector no puede acordar", acu([{"ref": k1, "src": "mantener"}], cl=c2).status_code == 403)
r = acu([{"ref": k1, "src": "manual", "pct": -15}, {"ref": k2, "src": "propuesta"}], "Reunión; comercial\nNelly ñ €")
check("acuerdos guardados", r.status_code == 200 and r.json["n"] == 2, r.json)
a1 = c.get(f"/api/desviacion?ref={k1}").json["rows"][0]
check("acuerdo vigente aplicado a la previsión", a1["estado"] == "acordado" and a1["ca"] == -0.15 and a1["src"] == "manual", a1)
check("acuerdo de la propuesta", c.get(f"/api/desviacion?ref={k2}").json["rows"][0]["ca"] == prop[1]["cp"])
txt = c.get("/api/desviacion/acuerdos.csv").data.decode("latin-1")
lin = txt.split("\r\n")
check("CSV de acuerdos: cabecera con los meses", lin[0].startswith("Referencia;Mandante;Marca;Versión;Corrección %;Motivo;") and len(lin[0].split(";")) == 18, lin[0])
fila = next(x for x in lin if x.startswith(k1 + ";"))
check("CSV de acuerdos: % con coma y motivo limpio", fila.split(";")[4] == "-15" and ";" not in fila.split(";")[5] and len(fila.split(";")) == 18, fila)
h = c.get(f"/api/desviacion/{k1}/historial").json
check("historial del acuerdo con versión y autor", len(h) == 1 and h[0]["version"] == ds["meta"]["version"] and h[0]["by"], h)
con = sqlite3.connect(A.DB_PATH)
con.execute("UPDATE prev_dec SET version='2026Q3' WHERE ref=?", (k2,))
con.commit()
con.close()
check("acuerdo de otra versión no es vigente", c.get(f"/api/desviacion?ref={k2}").json["rows"][0]["estado"] != "acordado")

# Nombres cortos de las líneas ("piedra Rosetta")
ln_ = c.get("/api/lineas").json
check("líneas: códigos de la carga con su nombre ABAS y, al principio, el código como nombre", ln_ and all(x["nombre"] == x["codigo"] for x in ln_) and any(x["abas"] for x in ln_), ln_[:2])
cod = next(x["codigo"] for x in ln_ if x["codigo"] != "—")
check("líneas: un lector no puede renombrar", c2.put("/api/lineas", json={"nombres": {cod: "Corta"}}, headers=H).status_code == 403)
check("líneas: nombre demasiado largo se rechaza", c.put("/api/lineas", json={"nombres": {cod: "x" * 41}}, headers=H).status_code == 400)
check("líneas: código que no está en la carga se rechaza", c.put("/api/lineas", json={"nombres": {"NOEXISTE": "x"}}, headers=H).status_code == 400)
check("líneas: renombrar", c.put("/api/lineas", json={"nombres": {cod: "  Espada 1 "}}, headers=H).status_code == 200
      and next(x for x in c.get("/api/lineas").json if x["codigo"] == cod)["nombre"] == "Espada 1")
check("líneas: el dataset trae los nombres", c.get("/api/dataset").json.get("alias", {}).get(cod) == "Espada 1")
c.put("/api/lineas", json={"nombres": {cod: ""}}, headers=H)
# Lote mínimo por línea: sube el lote calculado de sus referencias; vacío lo quita
_pl = next((p for p in c.get("/api/parametros").json["rows"] if p["ln"] and p["lc"] > 0 and not p["d"]), None)
if _pl:
    _lput = lambda v, cl=c: cl.put("/api/lineas", json={"lmin": {_pl["ln"]: v}}, headers=H)  # noqa: E731
    _pr = lambda: next(p for p in c.get("/api/parametros").json["rows"] if p["k"] == _pl["k"])  # noqa: E731
    check("lote mínimo: un lector no puede", _lput(9_000_000, c2).status_code == 403)
    check("lote mínimo no válido se rechaza", all(_lput(v).status_code == 400 for v in (0, -5, 12.5, "5000", True)))
    r = _lput(9_000_000)
    p = _pr()
    check("lote mínimo: guardado y aplicado al lote propuesto", r.status_code == 200 and next(x for x in r.json if x["codigo"] == _pl["ln"])["lmin"] == 9_000_000
          and p["lc"] == 9_000_000 and p["ltp"] == 9_000_000 and p["lms"], (r.status_code, p["lc"], p["ltp"]))
    _xo = openpyxl.load_workbook(io.BytesIO(c.get("/api/parametros.xlsx").data)).active
    _ho = [c_.value for c_ in _xo[1]]
    check("lote mínimo: el Excel lo indica", any(x[0] == _pl["k"] and "mínimo de la línea" in (dict(zip(_ho, x))["Observaciones"] or "")
                                                  for x in _xo.iter_rows(min_row=2, values_only=True)))
    _lput(None)
    check("lote mínimo: vacío lo quita", _pr()["lc"] == _pl["lc"] and not _pr()["lms"], _pr()["lc"])
check("líneas: vacío vuelve al código", next(x for x in c.get("/api/lineas").json if x["codigo"] == cod)["nombre"] == cod and cod not in c.get("/api/dataset").json.get("alias", {}))

# Capacidad por línea: valores iniciales del Excel, edición y configuración general
dsj = c.get("/api/dataset").json
check("capacidad: sembrada con el Excel", dsj.get("cap", {}).get("AER-01") == {"vmax": 4200, "oee": 0.601, "turnos": 1}, dsj.get("cap", {}).get("AER-01"))
check("capacidad: /api/lineas trae los parámetros", any(x["cap"] for x in c.get("/api/lineas").json))
check("capacidad: un lector no puede cambiarla", c2.put("/api/lineas", json={"cap": {cod: {"vmax": 1000, "oee": 0.5, "turnos": 2}}}, headers=H).status_code == 403)
check("capacidad: OEE fuera de rango se rechaza", c.put("/api/lineas", json={"cap": {cod: {"vmax": 1000, "oee": 50, "turnos": 1}}}, headers=H).status_code == 400)
check("capacidad: más de 3 turnos se rechaza", c.put("/api/lineas", json={"cap": {cod: {"vmax": 1000, "oee": 0.5, "turnos": 4}}}, headers=H).status_code == 400)
check("capacidad: guardar", c.put("/api/lineas", json={"cap": {cod: {"vmax": 1000, "oee": 0.5, "turnos": 2}}}, headers=H).status_code == 200
      and c.get("/api/dataset").json["cap"][cod] == {"vmax": 1000, "oee": 0.5, "turnos": 2})
c.put("/api/lineas", json={"cap": {cod: None}}, headers=H)
check("capacidad: vacío borra", cod not in c.get("/api/dataset").json["cap"])
# Turnos por mes: excepciones a los turnos de la línea
_tu = lambda body, cl=c, ln="LA3": cl.put(f"/api/lineas/{ln}/turnos", json=body, headers=H)  # noqa: E731
check("turnos por mes: un lector no puede cambiarlos", _tu({"meses": {"2026-11": 2}}, c2).status_code == 403)
check("turnos por mes: más de 3 se rechaza", _tu({"meses": {"2026-11": 4}}).status_code == 400)
check("turnos por mes: mes mal escrito se rechaza", _tu({"meses": {"2026-13": 2}}).status_code == 400)
check("turnos por mes: línea sin capacidad se rechaza", _tu({"meses": {"2026-11": 2}}, ln=cod).status_code == 400)
r = _tu({"meses": {"2026-11": 3, "2026-12": 2.5, "2027-01": 0}})
check("turnos por mes: guardar", r.status_code == 200 and c.get("/api/dataset").json["cap"]["LA3"]["meses"] == {"2026-11": 3, "2026-12": 2.5, "2027-01": 0},
      c.get("/api/dataset").json["cap"]["LA3"].get("meses"))
_tu({"meses": {"2026-11": None}})
check("turnos por mes: vacío vuelve a los de la línea", "2026-11" not in c.get("/api/dataset").json["cap"]["LA3"]["meses"])
cfgc = c.get("/api/config").json["capacidad"]
check("capacidad: configuración por defecto", cfgc["horas_turno"] == 7.75 and cfgc["holgura"] == 0.2 and cfgc["dias"]["2027-01"] == 18, cfgc["dias"].get("2027-01"))
check("capacidad: holgura fuera de rango se rechaza", c.put("/api/config", json={"capacidad": {"horas_turno": 8, "holgura": 2, "dias": {}}}, headers=H).status_code == 400)
check("capacidad: mes mal escrito se rechaza", c.put("/api/config", json={"capacidad": {"horas_turno": 8, "holgura": 0.1, "dias": {"2027-13": 20}}}, headers=H).status_code == 400)
cfgc = c.put("/api/config", json={"capacidad": {"horas_turno": 8, "holgura": 0.1, "dias": {"2027-01": 19}}}, headers=H).json["capacidad"]
check("capacidad: guardar configuración conserva los demás meses", cfgc["horas_turno"] == 8 and cfgc["dias"]["2027-01"] == 19 and cfgc["dias"]["2027-02"] == 20)
cfgc = c.put("/api/config", json={"capacidad": {"horas_turno": 8, "holgura": 0.1, "dias": {"2027-01": None}}}, headers=H).json["capacidad"]
check("capacidad: mes vacío vuelve a lunes a viernes", "2027-01" not in cfgc["dias"] and cfgc["dias"]["2027-02"] == 20)
check("capacidad: el dataset trae el stock mínimo para la carga", any("ss" in r for r in dsj["refs"]))

# Límite de intentos: la IP de X-Forwarded-For no cuenta si no hay proxy de confianza
c3 = A.app.test_client()
for i in range(8):
    c3.post("/api/login", json={"username": "admin", "password": "mal"}, headers={**H, "X-Forwarded-For": f"10.0.0.{i}"})
r = c3.post("/api/login", json={"username": "admin", "password": "mal"}, headers={**H, "X-Forwarded-For": "10.9.9.9"})
check("cambiar X-Forwarded-For no evita el bloqueo", r.status_code == 429)

# Paridad core.py / static/core.js
node = shutil.which("node")
if node:
    exc = c.get("/api/me").json["config"]["exceso"]
    out = {f"{e}{p}{h}": {r["k"]: core.evaluate(r, h, e, p, exc)["sem"] for r in ds["refs"]}
           for e in ("OF", "OFPF", "ALL") for p in ("T", "C") for h in (1, 3, 6)}
    with tempfile.TemporaryDirectory() as tmp:
        fj = os.path.join(tmp, "d.json")
        json.dump({"refs": ds["refs"], "py": out, "exc": exc}, open(fj, "w"))
        js = ("global.window={};require(process.argv[1]);const d=require(process.argv[2]);let n=0;"
              "for(const e of ['OF','OFPF','ALL'])for(const p of ['T','C'])for(const h of [1,3,6])for(const r of d.refs)"
              "if(window.Cob.evaluate(r,{horizonte:h,escenario:e,prevision:p,exceso:d.exc}).sem!==d.py[e+p+h][r.k])n++;console.log(n)")
        diff = subprocess.check_output([node, "-e", js, os.path.join(ROOT, "static", "core.js"), fj], text=True).strip()
    check("core.js evalúa igual que core.py", diff == "0", f"{diff} diferencias")
    check("hay referencias en exceso con los datos reales", sum(v == "exceso" for v in out["ALLT3"].values()) > 0, sum(v == "exceso" for v in out["ALLT3"].values()))
else:
    print("--   sin Node.js: no se comprueba la paridad con core.js")

print("\nTodo correcto" if not check.fails else f"\n{check.fails} comprobaciones fallidas")
sys.exit(1 if check.fails else 0)
